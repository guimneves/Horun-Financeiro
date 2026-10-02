"""Importação do orçamento a partir da planilha de acompanhamento (aba
"Saldo por Item") — pedido do usuário: cadastrar os ~190 itens à mão antes
de ler as pastas do drive era o passo mais demorado.

Formato da aba (o da planilha real do laboratório, sem valores aqui):

  Elemento de Despesa: Material de Consumo        <- título do elemento
  MATERIAL DE CONSUMO NACIONAL                    <- título da seção (opcional)
  Nº | Descrição do item | Finalidade/ Justificativa | ... | Valor (R$) | Rendimentos | ...
  1  | Reagente X        | Para ...                  | ... | 1500       | 0           | ...
  2  |                   |                           |     | 0          |             |      <- linha vazia, ignorada
  Total nacional = ...                                                                     <- fim da seção

O cabeçalho muda de seção para seção: equipamentos, diárias, outros bens e
protótipos têm "V. unitário" e "Quant."; obras, passagens, material de
consumo, serviços e outras despesas só "Valor (R$)" (aí o item entra com
quantidade 1 e valor unitário = valor). Por isso as colunas são achadas
pelo título, nunca pela posição. A Equipe Executora fica de fora: é
pessoal, com tela própria.

Nada é gravado aqui: `parse_budget_sheet` só devolve o plano; quem grava é
a rota (numa revisão NOVA, em rascunho, para conferir antes de ativar).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from app.core.money import round_money
from app.services.drive_scan import fold

SHEET_NAME = "Saldo por Item"

# (trecho do título, sem acento e minúsculo) -> categoria. Ordem importa:
# o primeiro que casar vale ("importad" antes do nacional de cada grupo).
_SECTION_RULES: list[tuple[tuple[str, ...], str | None]] = [
    (("equipe executora",), None),  # pessoal: fora da importação
    (("equipamento", "importad"), "equip_importado"),
    (("equipamento",), "equip_nacional"),
    (("obras",), "obras_instalacoes"),
    (("passage",), "passagens"),
    (("diaria",), "diarias"),
    (("ajuda de custo",), "diarias"),
    (("material de consumo", "importad"), "material_consumo_importado"),
    (("material de consumo",), "material_consumo_nacional"),
    (("servico",), "servicos_terceiros"),
    (("outros bens",), "outros_bens_direitos"),
    (("prototipo", "importad"), "prototipo_importado"),
    (("prototipo",), "prototipo_nacional"),
    (("despesas acessorias",), "outras_despesas"),
    (("outras despesas",), "outras_despesas"),
]


class BudgetImportError(ValueError):
    """Planilha que não dá para ler (sem a aba, arquivo inválido...)."""


@dataclass
class ImportedItem:
    category: str
    item_number: int
    description: str
    justification: str
    unit_value: Decimal
    planned_quantity: Decimal
    planned_value: Decimal
    yield_amount: Decimal
    sheet_row: int


@dataclass
class BudgetSheetResult:
    items: list[ImportedItem] = field(default_factory=list)
    skipped_sections: list[str] = field(default_factory=list)  # ex. "Equipe Executora (pessoal — tela própria)"
    warnings: list[str] = field(default_factory=list)


def _category_for(*titles: str | None) -> tuple[bool, str | None]:
    """(reconhecido, categoria). Tenta o título da seção e, se não der, o do
    elemento. Categoria None + reconhecido = seção que fica de fora (pessoal)."""
    for title in titles:
        if not title:
            continue
        folded = fold(title)
        for needles, category in _SECTION_RULES:
            if all(n in folded for n in needles):
                return True, category
    return False, None


def _is_header(row: tuple) -> bool:
    first = fold(str(row[0])) if row and row[0] is not None else ""
    second = fold(str(row[1])) if len(row) > 1 and row[1] is not None else ""
    return first in ("no", "nº", "n°", "n") or (first.startswith("n") and len(first) <= 3 and "descri" in second)


def _columns(header: tuple) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, cell in enumerate(header):
        if cell is None:
            continue
        name = fold(str(cell))
        if i == 0:
            cols["number"] = 0
        elif "descri" in name and "description" not in cols:
            cols["description"] = i
        elif "finalidade" in name or "justificativa" in name:
            cols.setdefault("justification", i)
        elif ("unitario" in name or name.startswith("v. unit")) and "unit" not in cols:
            cols["unit"] = i
        elif name.startswith("quant") and "disponivel" not in name and "utilizada" not in name:
            cols.setdefault("quantity", i)
        elif name.startswith("valor (r$)") or name == "valor":
            cols.setdefault("value", i)
        elif name.startswith("rendimento"):
            cols.setdefault("yield", i)
    return cols


def _number(raw, label: str, row_number: int, warnings: list[str]) -> Decimal:
    if raw is None or raw == "":
        return Decimal("0")
    try:
        value = Decimal(str(raw).replace(",", ".")) if isinstance(raw, str) else Decimal(str(raw))
    except InvalidOperation:
        warnings.append(f"Linha {row_number}: {label} {raw!r} não é número — usado 0.")
        return Decimal("0")
    if value < 0:
        warnings.append(f"Linha {row_number}: {label} negativo ({value}) — usado 0.")
        return Decimal("0")
    return value


def _item_number(raw) -> int | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)) and float(raw).is_integer():
        return int(raw)
    match = re.fullmatch(r"\s*(\d+)\s*", str(raw))
    return int(match.group(1)) if match else None


def parse_budget_sheet(content: bytes) -> BudgetSheetResult:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise BudgetImportError("Instale o extra de importação (pip install -e '.[import]').") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # zip corrompido, não é xlsx...
        raise BudgetImportError(f"Não foi possível abrir a planilha: {exc}") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise BudgetImportError(f'A planilha não tem a aba "{SHEET_NAME}".')
        rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()

    result = BudgetSheetResult()
    seen: set[tuple[str, int]] = set()
    element_title: str | None = None
    section_title: str | None = None
    cols: dict[str, int] | None = None
    category: str | None = None
    in_section = False

    for row_number, row in enumerate(rows, start=1):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        first = row[0]
        first_text = "" if first is None else str(first).strip()

        if _is_header(row):
            known, category = _category_for(section_title, element_title)
            cols = _columns(row)
            in_section = True
            if not known:
                title = section_title or element_title or "(sem título)"
                result.warnings.append(f'Linha {row_number}: seção "{title}" não corresponde a nenhuma categoria — ignorada.')
                category = None
            elif category is None:
                label = "Equipe Executora (pessoal — cadastrada na tela de Pessoal)"
                if label not in result.skipped_sections:
                    result.skipped_sections.append(label)
            if category is not None and (
                cols.get("description") is None or (cols.get("value") is None and cols.get("unit") is None)
            ):
                result.warnings.append(f"Linha {row_number}: cabeçalho sem 'Descrição' ou 'Valor' — seção ignorada.")
                category = None
            continue

        number = _item_number(first)
        if number is None and in_section and re.fullmatch(r"\d+[.,]\d+", first_text):
            # subitem ("1.1"): o orçamento só tem número inteiro por item.
            # Não encerra a seção (antes, tudo depois dele se perdia).
            if category is not None and cols is not None:
                desc = row[cols["description"]] if cols["description"] < len(row) else None
                index = cols.get("value")
                raw_value = row[index] if index is not None and index < len(row) else None
                result.warnings.append(
                    f"Linha {row_number}: subitem {first_text} ({str(desc or '').strip()[:60]}, Valor {raw_value}) "
                    "não foi importado — o orçamento só aceita número inteiro; lance-o à mão num item."
                )
            continue
        if number is None:
            # título de elemento/seção, ou linha de total (fim da seção)
            folded = fold(first_text)
            in_section = False
            if folded.startswith("elemento de despesa"):
                element_title, section_title = first_text, None
            elif first_text and not folded.startswith(("total", "valor total")):
                section_title = first_text
            continue

        if not in_section or category is None or cols is None:
            continue
        description_raw = row[cols["description"]] if cols["description"] < len(row) else None
        description = "" if description_raw is None else str(description_raw).strip()
        if not description:
            continue  # linha numerada vazia (a planilha reserva até 40 por seção)

        def cell(key: str):
            index = cols.get(key)
            return row[index] if index is not None and index < len(row) else None

        justification = "" if cell("justification") is None else str(cell("justification")).strip()
        value = round_money(_number(cell("value"), "Valor", row_number, result.warnings))
        if cols.get("unit") is not None and cols.get("quantity") is not None:
            unit = round_money(_number(cell("unit"), "V. unitário", row_number, result.warnings))
            quantity = _number(cell("quantity"), "Quant.", row_number, result.warnings).quantize(Decimal("0.01"))
            planned = round_money(unit * quantity)
            if cols.get("value") is not None and abs(planned - value) > Decimal("0.01"):
                result.warnings.append(
                    f"Linha {row_number} (item {number}): V. unitário × Quant. = {planned}, mas a planilha diz "
                    f"Valor {value} — foi usado V. unitário × Quant."
                )
        else:
            unit, quantity, planned = value, Decimal("1.00"), value

        key = (category, number)
        if key in seen:
            result.warnings.append(f"Linha {row_number}: item {number} repetido na mesma categoria — ignorado.")
            continue
        seen.add(key)
        result.items.append(ImportedItem(
            category=category, item_number=number, description=description, justification=justification,
            unit_value=unit, planned_quantity=quantity, planned_value=planned,
            yield_amount=round_money(_number(cell("yield"), "Rendimentos", row_number, result.warnings)),
            sheet_row=row_number,
        ))
    if not result.items:
        result.warnings.append(f'Nenhum item encontrado na aba "{SHEET_NAME}".')
    return result

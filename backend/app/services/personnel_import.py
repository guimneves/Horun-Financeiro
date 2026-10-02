"""Importação de quem ocupa cada vaga da Equipe Executora, a partir da aba
"Equipe Executora" da planilha de acompanhamento (pedido do usuário). As
vagas em si (itens do orçamento) vêm de `budget_import` — importe o
orçamento antes.

Formato da aba (o da planilha real, sem dados aqui), cabeçalho em 2 linhas:

  Item | Membro | Profissional |          | Início                 | Fim                    | Período | Valor | Valor total
       |        |              | Situação | Referência | Pagto     | Referência | Pagto     |         |       |
  1    | Vaga X | Fulana       | Ativo    | 01/05/2024 | 01/06/2024| 17/09/2026 | ...       | 29      | 781.44| 22661.76

- "Item" = nº da vaga no orçamento (pode juntar vagas: "11, 22, 23, 24" ->
  vaga 11, como no orçamento); a mesma vaga aparece mais de uma vez quando
  a pessoa mudou.
- "Membro" = o nome da vaga; "Profissional" = a pessoa.
- Datas de REFERÊNCIA (mês trabalhado), não de pagamento: é o que
  `services/accrual.py` conta (meses de calendário, mês final incluído).
- Ativo: a planilha põe no "Fim" a data do dia em que foi atualizada; aqui
  a atribuição fica sem fim e o acúmulo vai até hoje.

Nada é gravado aqui — `parse_personnel_sheet` só devolve o plano.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.core.money import round_money
from app.services.budget_import import BudgetImportError
from app.services.drive_scan import fold

SHEET_NAME = "Equipe Executora"


@dataclass
class ImportedAssignment:
    item_number: int
    role_title: str
    person_name: str
    status: str  # ativo | encerrado
    start_date: date
    end_date: date | None
    monthly_rate: Decimal
    sheet_value: Decimal | None  # "Valor total" da planilha (realizado dela)
    sheet_row: int


@dataclass
class PersonnelSheetResult:
    assignments: list[ImportedAssignment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _as_date(raw) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _money(raw) -> Decimal | None:
    if raw is None or raw == "":
        return None
    try:
        return round_money(Decimal(str(raw).replace(",", ".")) if isinstance(raw, str) else Decimal(str(raw)))
    except InvalidOperation:
        return None


def _item_number(raw) -> int | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)) and float(raw).is_integer():
        return int(raw)
    numbers = re.findall(r"\d+", str(raw))
    return int(numbers[0]) if numbers and re.fullmatch(r"\s*\d+(\s*,\s*\d+)*\s*", str(raw)) else None


def _find_columns(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """(linha onde os dados começam, colunas). Procura o cabeçalho nas
    primeiras linhas: "Item"/"Membro"/"Profissional" e, na linha de baixo,
    "Situação" e as duas "Referência" (início e fim)."""
    for index, row in enumerate(rows[:15]):
        names = [fold(str(c)) if c is not None else "" for c in row]
        if "item" in names and "membro" in names and "profissional" in names:
            cols = {"item": names.index("item"), "role": names.index("membro"), "person": names.index("profissional")}
            for i, name in enumerate(names):
                if name == "inicio":
                    cols["start_group"] = i
                elif name == "fim":
                    cols["end_group"] = i
                elif name == "valor":
                    cols["rate"] = i
                elif name == "valor total":
                    cols["total"] = i
            below = [fold(str(c)) if c is not None else "" for c in (rows[index + 1] if index + 1 < len(rows) else ())]
            if "situacao" in below:
                cols["status"] = below.index("situacao")
            refs = [i for i, name in enumerate(below) if name == "referencia"]
            # a "Referência" de cada grupo fica na coluna do título (Início/Fim) ou logo depois
            for key, group in (("start", "start_group"), ("end", "end_group")):
                if group in cols:
                    after = [i for i in refs if i >= cols[group]]
                    if after:
                        cols[key] = after[0]
            missing = [k for k in ("status", "start", "rate") if k not in cols]
            if missing:
                raise BudgetImportError(
                    f'Aba "{SHEET_NAME}": não achei as colunas {", ".join(missing)} no cabeçalho.'
                )
            return index + 2, cols
    raise BudgetImportError(f'Aba "{SHEET_NAME}": não achei o cabeçalho (Item / Membro / Profissional).')


def parse_personnel_sheet(content: bytes) -> PersonnelSheetResult:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise BudgetImportError("Instale o extra de importação (pip install -e '.[import]').") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise BudgetImportError(f"Não foi possível abrir a planilha: {exc}") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise BudgetImportError(f'A planilha não tem a aba "{SHEET_NAME}".')
        rows = list(workbook[SHEET_NAME].iter_rows(max_col=40, values_only=True))
    finally:
        workbook.close()

    start, cols = _find_columns(rows)
    result = PersonnelSheetResult()

    def cell(row: tuple, key: str):
        index = cols.get(key)
        return row[index] if index is not None and index < len(row) else None

    for row_number, row in enumerate(rows[start:], start=start + 1):
        person = cell(row, "person")
        person = "" if person is None else str(person).strip()
        item_raw = cell(row, "item")
        if not person or fold(person) in ("#n/a", "0") or item_raw in (None, "", 0):
            continue  # linha reservada vazia (fórmula sem pessoa)
        number = _item_number(item_raw)
        if number is None:
            result.warnings.append(f"Linha {row_number}: item {item_raw!r} de {person} não é um nº de vaga — não importado.")
            continue
        status_text = fold(str(cell(row, "status") or ""))
        if status_text.startswith("ativ"):
            status = "ativo"
        elif status_text.startswith("encerr"):
            status = "encerrado"
        else:
            result.warnings.append(f"Linha {row_number}: situação {cell(row, 'status')!r} de {person} não é Ativo/Encerrado — não importado.")
            continue
        start_date = _as_date(cell(row, "start"))
        if start_date is None:
            result.warnings.append(f"Linha {row_number}: {person} sem data de início — não importado.")
            continue
        end_date = _as_date(cell(row, "end")) if status == "encerrado" else None
        if status == "encerrado" and end_date is None:
            result.warnings.append(f"Linha {row_number}: {person} encerrado sem data de fim — não importado.")
            continue
        rate = _money(cell(row, "rate"))
        if rate is None or rate <= 0:
            result.warnings.append(f"Linha {row_number}: {person} sem valor mensal — não importado.")
            continue
        role = cell(row, "role")
        result.assignments.append(ImportedAssignment(
            item_number=number,
            role_title=("" if role is None else str(role).strip()) or f"Vaga {number}",
            person_name=person,
            status=status,
            start_date=start_date,
            end_date=end_date,
            monthly_rate=rate,
            sheet_value=_money(cell(row, "total")),
            sheet_row=row_number,
        ))
    if not result.assignments:
        result.warnings.append(f'Nenhuma pessoa encontrada na aba "{SHEET_NAME}".')
    return result

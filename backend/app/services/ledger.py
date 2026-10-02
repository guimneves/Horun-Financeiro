"""Leitura dos lançamentos da planilha de acompanhamento de saldo — as abas
por categoria ("Equip. Nacional", "Material de Consumo"...), onde cada linha é
um processo com favorecido, descrição, valor e nº de processo COPPETEC.

As pastas dizem QUAIS processos existem; esta planilha diz QUANTO cada um
custou e quem foi o favorecido. O vínculo é o nº de processo (normalizado).

As colunas são achadas pelo TÍTULO no cabeçalho (linha 2), não por posição:
umas abas têm a coluna "Quantidade" e outras não, o que desloca as demais.
Requer `openpyxl` (extra opcional `import`).
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from app.core.drive import fs_path
from app.core.process_number import normalize_process_number

# Nome da aba -> chave da categoria.
SHEET_CATEGORIES: dict[str, str] = {
    "Equip. Nacional": "equip_nacional",
    "Equip. Importado": "equip_importado",
    "Obras e Instalações": "obras_instalacoes",
    "Passagens": "passagens",
    "Diárias": "diarias",
    "Material de Consumo": "material_consumo_nacional",
    "Material de Consumo Importado": "material_consumo_importado",
    "Serv. Terceiros": "servicos_terceiros",
    "Outros Bens e Direitos": "outros_bens_direitos",
    "Protótipo Nacional": "prototipo_nacional",
    "Protótipo Importado": "prototipo_importado",
    "Outras Despesas": "outras_despesas",
}

_NUMBER_RE = re.compile(r"^\d{4}\D+\d+$")


class LedgerError(Exception):
    pass


@dataclass
class LedgerEntry:
    process_number: str
    category: str
    item_number: int | None
    vendor: str | None
    description: str | None
    quantity: Decimal
    value: Decimal
    lines: int = 1  # quantas linhas da planilha foram somadas neste processo


@dataclass
class LedgerResult:
    entries: dict[str, LedgerEntry]
    skipped: list[str]  # linhas com nº de processo que não é um nº válido, ou sem valor


def _as_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _find_columns(header: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(header):
        text = str(cell or "").casefold()
        if "processo" in text:
            columns["process"] = index
        elif "nº do item" in text or "n° do item" in text or "no do item" in text:
            columns["item"] = index
        elif "quantidade" in text:
            columns["quantity"] = index
        elif "favorecido" in text:
            columns["vendor"] = index
        elif "descri" in text:
            columns["description"] = index
        elif text.strip() == "valor":
            columns["value"] = index
    return columns


def read_ledger(source: str | bytes) -> LedgerResult:
    """`source`: caminho de um .xlsx no disco deste servidor, ou o conteúdo
    dele (como chega do Horun Agent, no modo agente)."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise LedgerError("Instale o extra de importação (pip install -e '.[import]') para ler a planilha.") from exc

    if isinstance(source, bytes):
        target = io.BytesIO(source)
    else:
        if not os.path.isfile(fs_path(source)):
            raise LedgerError(f"Planilha não encontrada: {source}")
        target = fs_path(source)
    try:
        workbook = openpyxl.load_workbook(target, data_only=True, read_only=True)
    except Exception as exc:  # zip corrompido, não é xlsx de verdade...
        raise LedgerError(f"Não foi possível abrir a planilha: {exc}") from exc
    entries: dict[str, LedgerEntry] = {}
    skipped: list[str] = []
    try:
        for sheet_name, category in SHEET_CATEGORIES.items():
            if sheet_name not in workbook.sheetnames:
                continue
            rows = workbook[sheet_name].iter_rows(min_row=2, max_col=12, values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            columns = _find_columns(header)
            if "process" not in columns or "value" not in columns:
                skipped.append(f"{sheet_name}: cabeçalho sem 'Valor' ou 'Nº de Processo' — aba ignorada.")
                continue
            for line_number, row in enumerate(rows, start=3):
                raw_number = row[columns["process"]] if columns["process"] < len(row) else None
                if raw_number in (None, ""):
                    continue
                value = _as_decimal(row[columns["value"]])
                if not _NUMBER_RE.match(str(raw_number).strip()):
                    amount = f" (valor {value})" if value is not None else ""
                    skipped.append(
                        f"{sheet_name} linha {line_number}: '{raw_number}' não parece um nº de processo{amount}."
                    )
                    continue
                if value is None:
                    skipped.append(f"{sheet_name} linha {line_number}: processo {raw_number} sem valor.")
                    continue
                number = normalize_process_number(str(raw_number))
                assert number is not None

                def cell(key: str) -> object:
                    index = columns.get(key)
                    return row[index] if index is not None and index < len(row) else None

                quantity = _as_decimal(cell("quantity")) or Decimal("1")
                item_raw = cell("item")
                item_number = int(item_raw) if isinstance(item_raw, (int, float)) else None
                if number in entries:
                    existing = entries[number]
                    existing.value += value
                    existing.lines += 1
                else:
                    entries[number] = LedgerEntry(
                        process_number=number,
                        category=category,
                        item_number=item_number,
                        vendor=(str(cell("vendor")).strip() or None) if cell("vendor") else None,
                        description=(str(cell("description")).strip() or None) if cell("description") else None,
                        quantity=quantity,
                        value=value,
                    )
    finally:
        workbook.close()
    return LedgerResult(entries=entries, skipped=skipped)

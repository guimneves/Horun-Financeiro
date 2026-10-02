"""Importar o orçamento da aba "Saldo por Item" da planilha de
acompanhamento. Planilha SINTÉTICA, montada aqui com a mesma estrutura da
real (repositório público: nada de dado real)."""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest

from app.services.budget_import import parse_budget_sheet
from tests.conftest import ADMIN, COLAB

HEADER_UNIT = ["Nº", "Descrição do item", "Finalidade/ Justificativa", "V. unitário", "Quant. Prevista", "Valor (R$)",
               "Rendimentos", "Quant. Disponível", "Valor Realizado (R$)", "Saldo do Item (R$)"]
HEADER_VALUE = ["Nº", "Descrição do item", "Finalidade/ Justificativa", "Valor (R$)", "Rendimentos",
                "Valor Realizado (R$)", "Saldo do Item (R$)"]
HEADER_PERSONNEL = ["Nº", "Membro da Equipe", "Profissional", "Tipo de Remuneração (3)", "Modalidade (4)",
                    "Período (em meses)", "Carga Horária Semanal", "Valor Mensal (R$)", "Encargos",
                    "Valor Total Mensal\n(R$)", "Valor Total\n(R$)", "Quantidade Utilizada"]


def _sheet() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Quadro Resumo"
    ws = wb.create_sheet("Saldo por Item")
    rows = [
        ["Elemento de Despesa: Equipamento e Material Permanente"],
        ["EQUIPAMENTO E MATERIAL PERMANENTE NACIONAL"],
        HEADER_UNIT,
        [1, "Computador", "Aquisição de dados", 5000, 2, 10000, 0, 0, 5000, 5000],
        [2, "Bomba", "Vácuo", 7000, 0, 0, 0, 0, 0, 0],  # previsto zero: entra mesmo assim
        [3, None, None, None, None, 0, 0, 0, 0, 0],  # linha reservada vazia
        ["Total Equipamento e Material Permanente Nacional", None, None, None, None, 10000],
        ["EQUIPAMENTO E MATERIAL PERMANENTE IMPORTADO"],
        HEADER_UNIT,
        [1, "Booster", "Pressão", 50000, 1, 49000, 0, 0, 0, 0],  # V.unit × Quant ≠ Valor
        ["VALOR TOTAL DO ELEMENTO DE DESPESA"],
        [None, None, None, None, None, None, "Rendimentos"],
        ["Elemento de Despesa: Equipe Executora"],
        HEADER_PERSONNEL,
        [1, "Vaga técnico A", "Técnico", "Ressarcimento - HH", "-", 24, 5, 500, 100, 600, 14400, 3],
        [2, "PD - Vaga B", "PD - Fulana", "Bolsa de pesquisa", "BOLSA - PÓS-DOUTORADO", 12, 40, 7000, 0, 7000, 84000, 2],
        [5, "Pessoa C", "Tec", "Remuneração Direta", "-", None, 40, 4000, 3000, 7000, 0, 0],  # valor antigo
        [6, "Pessoa C", "Tec", "Remuneração Direta", "-", None, 40, 4200, 3100, 7300, 0, 0],  # valor antigo
        [7, None, None, None, None, None, None, None, None, 0, 0, 0],
        ["5, 6", "Pessoa C", "Tec - C", "Remuneração Direta", "-", 10, 40, 4400, 3200, 7600, 76000, 1],
        ["VALOR TOTAL DO ELEMENTO DE DESPESA"],
        ["Elemento de Despesa: Equipe Executora"],  # segunda tabela (plano antigo): ignorada
        ["Nº", "Descrição do cargo (1)", "Finalidade/ Justificativa (2)", "Tipo de Remuneração (3)", "Valor Total"],
        [1, "Cargo antigo", "Algo", "Bolsa", 999],
        ["VALOR TOTAL DO ELEMENTO DE DESPESA"],
        ["Elemento de Despesa: Serviços de Terceiros"],
        ["OUTRAS DESPESAS COM SERVIÇOS DE TERCEIROS (Pessoa Jurídica)"],
        HEADER_VALUE,
        [1, "Manutenção", "Corretiva", 3000, 0],
        ["1.1", "Subitem", "Algo", 500, 0],  # antes, isto encerrava a seção
        [2, "Usinagem", "Peças", 1200.5, 0],
        [3, "Reparo pago com rendimento", "Rendimentos", 0, 800],
        ["Total Serviços de Terceiros =", None, None, 4700.5],
        ["Elemento de Despesa: Diárias (Pessoal Civil / Militar)"],
        ["Nº", "Descrição do item", "Finalidade/ Justificativa", "V. unitário", "Quant.", "Valor (R$)", "Rendimentos"],
        [1, "Diária nacional", "Campo", 300, 10, 3000, 0],
    ]
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_reads_every_section_by_its_own_header():
    result = parse_budget_sheet(_sheet())
    got = {(i.category, i.item_number): i for i in result.items}
    assert set(got) == {
        ("equip_nacional", 1), ("equip_nacional", 2), ("equip_importado", 1),
        ("equipe_executora", 1), ("equipe_executora", 2), ("equipe_executora", 5),
        ("servicos_terceiros", 1), ("servicos_terceiros", 2), ("servicos_terceiros", 3),
        ("diarias", 1),
    }
    pc = got[("equip_nacional", 1)]
    assert (pc.unit_value, pc.planned_quantity, pc.planned_value) == (Decimal("5000"), Decimal("2.00"), Decimal("10000"))
    assert pc.justification == "Aquisição de dados"
    # só "Valor (R$)": quantidade 1, unitário = valor
    usinagem = got[("servicos_terceiros", 2)]
    assert (usinagem.unit_value, usinagem.planned_quantity) == (Decimal("1200.50"), Decimal("1.00"))
    assert got[("servicos_terceiros", 3)].yield_amount == Decimal("800")
    assert got[("diarias", 1)].planned_value == Decimal("3000")


def test_personnel_vacancies_use_monthly_total_times_months():
    got = {(i.category, i.item_number): i for i in parse_budget_sheet(_sheet()).items}
    tech = got[("equipe_executora", 1)]
    assert (tech.unit_value, tech.planned_quantity, tech.planned_value) == (Decimal("600"), Decimal("24.00"), Decimal("14400"))
    assert tech.description == "Vaga técnico A" and tech.justification == "Técnico"
    assert "Ressarcimento - HH" in tech.note and "5 h/semana" in tech.note
    # vagas juntadas: fica a do 1º número, com o valor consolidado; as avulsas saem
    merged = got[("equipe_executora", 5)]
    assert merged.planned_value == Decimal("76000") and ("equipe_executora", 6) not in got


def test_parse_reports_what_it_could_not_import():
    result = parse_budget_sheet(_sheet())
    assert result.skipped_sections == ["Segunda tabela de Equipe Executora (linha 23) — usada só a da linha 14"]
    assert any("subitem 1.1" in w and "500" in w for w in result.warnings)
    assert any("item 1" in w and "49000" in w for w in result.warnings)  # V. unitário × Quant. ≠ Valor
    assert any("vagas 5, 6 juntadas" in w for w in result.warnings)
    assert len(result.warnings) == 3


def test_sheet_without_the_tab_is_refused():
    wb = openpyxl.Workbook()
    buffer = io.BytesIO()
    wb.save(buffer)
    from app.services.budget_import import BudgetImportError

    with pytest.raises(BudgetImportError, match="Saldo por Item"):
        parse_budget_sheet(buffer.getvalue())


def _project(client):
    return client.post("/projects", json={"code": "99.001", "name": "Projeto sintético"}, headers=ADMIN).json()


def _upload(client, url, headers=ADMIN, **form):
    return client.post(
        url, headers=headers, data=form,
        files={"file": ("planilha.xlsx", _sheet(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_preview_groups_by_category_and_writes_nothing(client):
    project = _project(client)
    r = _upload(client, f"/projects/{project['id']}/budget-import/preview")
    assert r.status_code == 200, r.text
    by_cat = {c["category"]: c for c in r.json()["categories"]}
    assert by_cat["servicos_terceiros"]["count"] == 3
    assert Decimal(by_cat["servicos_terceiros"]["planned_total"]) == Decimal("4200.50")
    assert Decimal(by_cat["servicos_terceiros"]["yield_total"]) == Decimal("800")
    assert client.get(f"/projects/{project['id']}/revisions", headers=ADMIN).json() == []


def test_import_creates_a_draft_revision_with_all_items(client):
    project = _project(client)
    r = _upload(client, f"/projects/{project['id']}/budget-import", label="Reformulação 3", effective_date="2026-01-01")
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["items_created"] == 10 and body["revision"]["status"] == "rascunho"
    revision_id = body["revision"]["id"]
    items = client.get(f"/projects/{project['id']}/revisions/{revision_id}/items", headers=ADMIN).json()
    assert {(i["category"], i["item_number"]) for i in items} >= {("equip_nacional", 1), ("diarias", 1)}

    # ativar continua sendo um passo à parte, depois de conferir
    assert client.post(f"/projects/{project['id']}/revisions/{revision_id}/activate", headers=ADMIN).status_code == 200


def test_reimport_reuses_positions_in_a_new_revision(client):
    project = _project(client)
    first = _upload(client, f"/projects/{project['id']}/budget-import", effective_date="2026-01-01").json()
    client.post(f"/projects/{project['id']}/revisions/{first['revision']['id']}/activate", headers=ADMIN)
    second = _upload(client, f"/projects/{project['id']}/budget-import", effective_date="2026-06-01")
    assert second.status_code == 201, second.text
    assert second.json()["revision"]["revision_number"] == first["revision"]["revision_number"] + 1


def test_only_the_coordinator_imports(client):
    project = _project(client)
    client.post(
        f"/projects/{project['id']}/members",
        json={"user_id": "u-colab", "username": "colaborador", "role": "colaborador"}, headers=ADMIN,
    )
    assert _upload(client, f"/projects/{project['id']}/budget-import/preview", headers=COLAB).status_code == 403


def test_not_an_xlsx_is_refused(client):
    project = _project(client)
    r = client.post(
        f"/projects/{project['id']}/budget-import/preview", headers=ADMIN,
        files={"file": ("planilha.csv", b"a,b", "text/csv")},
    )
    assert r.status_code == 422


def test_me_says_who_is_core_admin(client):
    assert client.get("/auth/me", headers=ADMIN).json()["is_core_admin"] is True
    assert client.get("/auth/me", headers=COLAB).json()["is_core_admin"] is False
